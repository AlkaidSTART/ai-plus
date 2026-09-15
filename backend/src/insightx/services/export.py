"""Export service for packaging engineering charter into a zip archive."""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, datetime
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from insightx.errors import ApiError
from insightx.models import Evidence, Report, Task, TaskItem
from insightx.schemas import TaskStatus


def _str_width(value: Any) -> int:
    """Calculate approximate character display width for auto-sizing columns."""
    text = str(value if value is not None else "")
    return sum(2 if ord(char) > 127 else 1 for char in text)


def _build_product_xlsx(
    task: Task,
    item: TaskItem,
    report: Report | None,
    evidences: list[Evidence],
) -> bytes:
    """Generate product spreadsheet containing overview, proposals, pain points, and evidence."""
    # ponytail: fixed-schema openpyxl tables; upgrade to dynamic report schemas if new data types added
    wb = openpyxl.Workbook()

    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="1E293B")
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    cell_font = Font(name="Microsoft YaHei", size=9)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    def style_headers(ws: Any, col_count: int) -> None:
        for col in range(1, col_count + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
            cell.border = thin_border

    def auto_fit_columns(ws: Any) -> None:
        for col in ws.columns:
            max_w = max(_str_width(c.value) for c in col)
            letter = get_column_letter(col[0].column)
            ws.column_dimensions[letter].width = min(max(max_w + 3, 12), 60)
            for c in col:
                if c.row != 1:
                    c.font = cell_font
                    c.border = thin_border
                    c.alignment = Alignment(vertical="center", wrap_text=True)

    # 1. Sheet: 产品概览
    ws_meta = wb.active
    ws_meta.title = "产品概览"
    ws_meta.append(["属性 / 指标", "对应取值"])
    style_headers(ws_meta, 2)

    sample_metrics = report.sample_metrics if report else (item.sample_metrics or {})
    raw_count = sample_metrics.get("raw_review_count", 0)
    valid_count = sample_metrics.get("valid_review_count", 0)
    excluded_count = sample_metrics.get("excluded_review_count", 0)
    avg_rating = sample_metrics.get("average_rating")
    neg_ratio = sample_metrics.get("negative_review_ratio")
    neg_ratio_str = f"{neg_ratio * 100:.1f}%" if neg_ratio is not None else "暂无"
    methodology = sample_metrics.get("methodology", "暂无")

    meta_rows = [
        ("任务编号", task.task_id),
        ("产品 ASIN", item.asin),
        ("数据质量", report.data_quality if report else (item.data_quality or "未评估")),
        ("财务评估状态", report.financial_state if report else "NOT_EVALUATED"),
        ("原始评论总数", raw_count),
        ("有效评论数", valid_count),
        ("过滤评论数", excluded_count),
        ("平均星级评分", f"{avg_rating:.2f}" if avg_rating is not None else "暂无"),
        ("负面评论占比", neg_ratio_str),
        ("提炼痛点总数", len(report.pain_points) if report else 0),
        ("落地改款建议数", len(report.proposals) if report else 0),
        ("分析口径与方法", methodology),
        ("导出时间", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S UTC")),
    ]
    for r in meta_rows:
        ws_meta.append(list(r))
    auto_fit_columns(ws_meta)

    # 2. Sheet: 改款建议
    ws_prop = wb.create_sheet(title="改款建议")
    prop_headers = [
        "建议编号",
        "建议分类",
        "建议标题",
        "改造方案说明",
        "关联痛点编号",
        "支撑证据编号",
        "开模费用",
        "改模周期",
        "运费节省",
    ]
    ws_prop.append(prop_headers)
    style_headers(ws_prop, len(prop_headers))

    proposals = report.proposals if report else []
    for prop in proposals:
        col_type = (
            "产品物理本体优化"
            if prop.get("column") == "PRODUCT_OPTIMIZATION"
            else "包装与履约优化"
        )
        ws_prop.append([
            prop.get("proposal_id", ""),
            col_type,
            prop.get("title", ""),
            prop.get("change_description", ""),
            ", ".join(prop.get("pain_point_ids", [])),
            ", ".join(prop.get("evidence_refs", [])),
            "未评估",
            "未评估",
            "未评估",
        ])
    auto_fit_columns(ws_prop)

    # 3. Sheet: 用户痛点
    ws_pain = wb.create_sheet(title="用户痛点")
    pain_headers = [
        "痛点编号",
        "痛点标签",
        "严重等级",
        "严重度得分 (1-5)",
        "严重度依据",
        "出现频次",
        "频次口径",
        "痛点摘要",
        "支撑证据编号",
    ]
    ws_pain.append(pain_headers)
    style_headers(ws_pain, len(pain_headers))

    pain_points = report.pain_points if report else []
    for p in pain_points:
        ws_pain.append([
            p.get("pain_point_id", ""),
            p.get("label", ""),
            p.get("severity_level", ""),
            p.get("severity_score", ""),
            p.get("severity_rationale", ""),
            p.get("actual_frequency", 0),
            p.get("frequency_methodology", ""),
            p.get("summary", ""),
            ", ".join(p.get("evidence_refs", [])),
        ])
    auto_fit_columns(ws_pain)

    # 4. Sheet: 原始证据链
    ws_evi = wb.create_sheet(title="原始证据链")
    evi_headers = [
        "证据编号",
        "评论来源引用",
        "星级评分",
        "原始评论摘录",
        "证据来源类型",
        "来源链接",
        "发布时间",
    ]
    ws_evi.append(evi_headers)
    style_headers(ws_evi, len(evi_headers))

    for ev in evidences:
        rating = ev.source_metadata.get("rating") if ev.source_metadata else None
        pub_time = (
            ev.published_at.strftime("%Y-%m-%d %H:%M:%S")
            if ev.published_at
            else ""
        )
        ws_evi.append([
            ev.evidence_id,
            ev.source_ref,
            rating if rating is not None else "暂无",
            ev.excerpt,
            ev.source_type,
            ev.source_url or "",
            pub_time,
        ])
    auto_fit_columns(ws_evi)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _build_proposals_docx(
    task: Task,
    items: list[TaskItem],
    reports_map: dict[str, Report],
    evidences_map: dict[str, dict[str, Evidence]],
) -> bytes:
    """Generate engineering task charter DOCX document."""
    doc = Document()

    # Document title
    title_p = doc.add_heading(level=0)
    title_run = title_p.add_run("工程改款任务书 (Engineering RFC)")
    title_run.font.bold = True
    title_run.font.size = Pt(20)

    sub_p = doc.add_paragraph("基于真实买家评论原声与证据链驱动的改款决策落地指南")
    sub_p.runs[0].font.size = Pt(11)
    sub_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    # Task Metadata Table
    meta_table = doc.add_table(rows=4, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = True

    meta_entries = [
        ("任务编号", task.task_id, "目标站点", f"{task.platform.upper()} - {task.marketplace}"),
        ("分析窗口", task.window_preset, "任务状态", task.status),
        ("涉及产品数", str(len(items)), "导出时间", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("落地原则", "100% 绑定原始真实评论证据", "工程数值状态", "未评估依据不捏造"),
    ]
    for r_idx, (k1, v1, k2, v2) in enumerate(meta_entries):
        row = meta_table.rows[r_idx].cells
        row[0].text = k1
        row[1].text = v1
        row[2].text = k2
        row[3].text = v2
        for c in row:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)

    doc.add_paragraph()

    # Chapter 1: 核心规范说明
    h1 = doc.add_heading("一、工程任务书执行规范", level=1)
    h1.runs[0].font.size = Pt(14)

    p_spec1 = doc.add_paragraph(
        "1. 证据驱动：所有产品改款决策严格基于买家真实负评与使用痛点提炼，"
        "每项建议均反查真实买家原声证据，杜绝模型伪造证据。"
    )
    p_spec1.runs[0].font.size = Pt(10.5)

    p_spec2 = doc.add_paragraph(
        "2. 双栏分离：改款建议划分为「产品物理本体优化」与「包装与履约优化」，"
        "分别交付结构模具工程师与供应链包装团队。"
    )
    p_spec2.runs[0].font.size = Pt(10.5)

    p_spec3 = doc.add_paragraph(
        "3. 工程边界说明：开模费用、改模周期、运费节省等工程数值仅在企业输入及评估依据齐全时呈现。"
        "依据系统 PRD 规范，当前阶段均明确标注为「未评估」，防止盲目立项。"
    )
    p_spec3.runs[0].font.size = Pt(10.5)

    # Chapter 2: 各产品改款决策清单
    h2 = doc.add_heading("二、各产品改款决策清单", level=1)
    h2.runs[0].font.size = Pt(14)

    for item in items:
        report = reports_map.get(item.item_id)
        evi_dict = evidences_map.get(item.item_id, {})

        h_prod = doc.add_heading(f"产品 ASIN: {item.asin}", level=2)
        h_prod.runs[0].font.size = Pt(12)

        sample_metrics = report.sample_metrics if report else (item.sample_metrics or {})
        raw_cnt = sample_metrics.get("raw_review_count", 0)
        valid_cnt = sample_metrics.get("valid_review_count", 0)
        neg_ratio = sample_metrics.get("negative_review_ratio")
        neg_str = f"{neg_ratio * 100:.1f}%" if neg_ratio is not None else "暂无"
        avg_rt = sample_metrics.get("average_rating")
        rt_str = f"{avg_rt:.2f}" if avg_rt is not None else "暂无"

        doc.add_paragraph(
            f"数据质量：{report.data_quality if report else (item.data_quality or '未评估')} | "
            f"财务状态：{report.financial_state if report else 'NOT_EVALUATED'} | "
            f"原始评论：{raw_cnt} 条 | 有效评论：{valid_cnt} 条 | "
            f"平均评分：{rt_str} | 负评占比：{neg_str}"
        )

        # 2.1 产品物理本体优化
        doc.add_heading("1. 产品物理本体优化建议 (Product Optimization)", level=3)
        prod_proposals = (
            [p for p in report.proposals if p.get("column") == "PRODUCT_OPTIMIZATION"]
            if report
            else []
        )
        if not prod_proposals:
            doc.add_paragraph("暂无本体优化建议（无证据的结论不作为有效建议）。")
        else:
            for p in prod_proposals:
                doc.add_paragraph(f"• 建议标题：{p.get('title', '')}", style="List Bullet")
                doc.add_paragraph(f"  改动方案：{p.get('change_description', '')}")
                doc.add_paragraph(f"  关联痛点：{', '.join(p.get('pain_point_ids', []))}")
                evi_refs = p.get("evidence_refs", [])
                if evi_refs:
                    doc.add_paragraph("  支撑证据摘录：")
                    for ref in evi_refs:
                        ev = evi_dict.get(ref)
                        if ev:
                            rating = ev.source_metadata.get("rating") if ev.source_metadata else None
                            doc.add_paragraph(
                                f"    - [{ref}] ⭐ {rating or '-'} 分: “{ev.excerpt}”"
                            )
                        else:
                            doc.add_paragraph(f"    - [{ref}]")

        # 2.2 包装与履约优化
        doc.add_heading("2. 包装与履约优化建议 (Packaging & Fulfillment)", level=3)
        pack_proposals = (
            [p for p in report.proposals if p.get("column") == "PACKAGING_FULFILLMENT_OPTIMIZATION"]
            if report
            else []
        )
        if not pack_proposals:
            doc.add_paragraph("暂无包装履约优化建议（无证据的结论不作为有效建议）。")
        else:
            for p in pack_proposals:
                doc.add_paragraph(f"• 建议标题：{p.get('title', '')}", style="List Bullet")
                doc.add_paragraph(f"  改动方案：{p.get('change_description', '')}")
                doc.add_paragraph(f"  关联痛点：{', '.join(p.get('pain_point_ids', []))}")
                evi_refs = p.get("evidence_refs", [])
                if evi_refs:
                    doc.add_paragraph("  支撑证据摘录：")
                    for ref in evi_refs:
                        ev = evi_dict.get(ref)
                        if ev:
                            rating = ev.source_metadata.get("rating") if ev.source_metadata else None
                            doc.add_paragraph(
                                f"    - [{ref}] ⭐ {rating or '-'} 分: “{ev.excerpt}”"
                            )
                        else:
                            doc.add_paragraph(f"    - [{ref}]")

        # 2.3 核心痛点清单
        doc.add_heading("3. 关联核心痛点清单", level=3)
        pain_points = report.pain_points if report else []
        if not pain_points:
            doc.add_paragraph("暂无提取到的有效痛点。")
        else:
            pain_tbl = doc.add_table(rows=1, cols=5)
            pain_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            pain_hdr = pain_tbl.rows[0].cells
            pain_hdr[0].text = "痛点编号"
            pain_hdr[1].text = "痛点标签"
            pain_hdr[2].text = "严重程度"
            pain_hdr[3].text = "频次"
            pain_hdr[4].text = "痛点简要描述"
            for pt in pain_points:
                row_cells = pain_tbl.add_row().cells
                row_cells[0].text = pt.get("pain_point_id", "")
                row_cells[1].text = pt.get("label", "")
                row_cells[2].text = f"{pt.get('severity_level', '')} ({pt.get('severity_score', '')}分)"
                row_cells[3].text = str(pt.get("actual_frequency", 0))
                row_cells[4].text = pt.get("summary", "")

        # 2.4 风险与预警
        doc.add_heading("4. 风险与预警提示", level=3)
        warnings = report.warnings if report else []
        if warnings:
            for w in warnings:
                doc.add_paragraph(f"• [{w.get('code', 'WARN')}] {w.get('message', '')}")
        else:
            doc.add_paragraph("• 样本分析过程中未发现高危阻断预警。")
        doc.add_paragraph("• 工程数值边界：模具公差、改模周期及运费核算当前标记为「未评估」。")

        doc.add_paragraph()

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def export_task_charter_zip(
    session: Session,
    *,
    tenant_id: str,
    task_id: str,
) -> tuple[bytes, str]:
    """Export the engineering task charter as a zip archive containing docx and per-product xlsx."""
    task = session.scalar(
        select(Task).where(Task.tenant_id == tenant_id, Task.task_id == task_id)
    )
    if task is None:
        raise ApiError(404, "TASK_NOT_FOUND", "The task was not found.")

    if task.status in (TaskStatus.QUEUED.value, TaskStatus.RUNNING.value):
        raise ApiError(
            409,
            "TASK_NOT_READY",
            "The task is still running. Please wait for completion before exporting.",
            retryable=True,
        )

    items = session.scalars(
        select(TaskItem)
        .where(TaskItem.tenant_id == tenant_id, TaskItem.task_id == task_id)
        .order_by(TaskItem.created_at.asc())
    ).all()

    # Load reports and evidences for all items
    reports = session.scalars(
        select(Report).where(Report.tenant_id == tenant_id, Report.task_id == task_id)
    ).all()
    reports_map: dict[str, Report] = {r.item_id: r for r in reports}

    evidences = session.scalars(
        select(Evidence).where(Evidence.tenant_id == tenant_id, Evidence.task_id == task_id)
    ).all()
    evidences_by_item: dict[str, list[Evidence]] = {}
    evidences_by_ref: dict[str, dict[str, Evidence]] = {}
    for ev in evidences:
        evidences_by_item.setdefault(ev.item_id, []).append(ev)
        evidences_by_ref.setdefault(ev.item_id, {})[ev.source_ref] = ev

    # Build DOCX charter recommendations
    docx_bytes = _build_proposals_docx(
        task=task,
        items=list(items),
        reports_map=reports_map,
        evidences_map=evidences_by_ref,
    )

    # Build ZIP package
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("工程任务书_建议.docx", docx_bytes)
        for item in items:
            item_report = reports_map.get(item.item_id)
            item_evidences = evidences_by_item.get(item.item_id, [])
            xlsx_bytes = _build_product_xlsx(
                task=task,
                item=item,
                report=item_report,
                evidences=item_evidences,
            )
            zf.writestr(f"{item.asin}.xlsx", xlsx_bytes)

    filename = f"工程任务书_{task_id}.zip"
    return zip_buffer.getvalue(), filename

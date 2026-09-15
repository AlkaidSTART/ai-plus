"""Tests for engineering charter export service and endpoint."""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, datetime
from typing import Any

import openpyxl  # type: ignore[import-untyped]
import pytest
from docx import Document
from sqlalchemy import BigInteger, create_engine
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from insightx.database import Base, get_session
from insightx.errors import ApiError
from insightx.main import create_app
from insightx.models import Evidence, Report, Task, TaskItem
from insightx.schemas import TaskStatus
from insightx.services.export import export_task_charter_zip


@compiles(JSONB, "sqlite")
def _compile_jsonb_sqlite(type_: Any, compiler: Any, **kw: Any) -> str:
    del type_, compiler, kw
    return "JSON"


@compiles(BigInteger, "sqlite")
def _compile_bigint_sqlite(type_: Any, compiler: Any, **kw: Any) -> str:
    del type_, compiler, kw
    return "INTEGER"


@pytest.fixture
def sqlite_session_factory():
    """Isolated in-memory SQLite session factory."""
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def _seed_completed_task(
    session_factory: sessionmaker[Session],
) -> tuple[str, str, str]:
    task_id = "tsk_export_test"
    item1_id = "itm_export_1"
    item2_id = "itm_export_2"
    asin1 = "B0FFWCNZGF"
    asin2 = "B0FFWGQHJP"
    now = datetime.now(UTC)

    with session_factory() as session:
        session.add(
            Task(
                task_id=task_id,
                tenant_id="dev-tenant",
                platform="amazon",
                marketplace="US",
                window_preset="6m",
                status=TaskStatus.COMPLETED.value,
                created_at=now,
                finished_at=now,
            )
        )
        session.add(
            TaskItem(
                item_id=item1_id,
                task_id=task_id,
                tenant_id="dev-tenant",
                asin=asin1,
                status=TaskStatus.COMPLETED.value,
                data_quality="SUFFICIENT",
                sample_metrics={
                    "raw_review_count": 120,
                    "valid_review_count": 100,
                    "excluded_review_count": 20,
                    "average_rating": 3.8,
                    "negative_review_ratio": 0.25,
                    "methodology": "Amazon US 评论分析",
                },
            )
        )
        session.add(
            TaskItem(
                item_id=item2_id,
                task_id=task_id,
                tenant_id="dev-tenant",
                asin=asin2,
                status=TaskStatus.COMPLETED.value,
                data_quality="NO_DATA",
                sample_metrics={
                    "raw_review_count": 0,
                    "valid_review_count": 0,
                    "excluded_review_count": 0,
                    "average_rating": None,
                    "negative_review_ratio": None,
                    "methodology": "Amazon US 评论分析",
                },
            )
        )
        session.add(
            Report(
                report_id="rpt_export_1",
                tenant_id="dev-tenant",
                task_id=task_id,
                item_id=item1_id,
                data_quality="SUFFICIENT",
                sample_metrics={
                    "raw_review_count": 120,
                    "valid_review_count": 100,
                    "excluded_review_count": 20,
                    "average_rating": 3.8,
                    "negative_review_ratio": 0.25,
                    "methodology": "Amazon US 评论分析",
                },
                financial_state="NOT_EVALUATED",
                pain_points=[
                    {
                        "pain_point_id": "pain_01",
                        "label": "卡扣松动易碎",
                        "severity_level": "CRITICAL",
                        "severity_score": 5,
                        "severity_rationale": "多次买家投诉断裂",
                        "actual_frequency": 35,
                        "frequency_methodology": "文本聚类统计",
                        "summary": "塑料卡扣在运输或安装中断裂",
                        "evidence_refs": ["rev_01"],
                    }
                ],
                proposals=[
                    {
                        "proposal_id": "prop_01",
                        "column": "PRODUCT_OPTIMIZATION",
                        "title": "加厚卡扣模具壁厚",
                        "change_description": "将卡扣受力部位壁厚增加至1.8mm",
                        "pain_point_ids": ["pain_01"],
                        "snapshot_ref": "snap_01",
                        "evidence_refs": ["rev_01"],
                    },
                    {
                        "proposal_id": "prop_02",
                        "column": "PACKAGING_FULFILLMENT_OPTIMIZATION",
                        "title": "加入高密度EPE防震珍珠棉",
                        "change_description": "包装四周增加15mm防震保护",
                        "pain_point_ids": ["pain_01"],
                        "snapshot_ref": "snap_02",
                        "evidence_refs": ["rev_01"],
                    },
                ],
                warnings=[
                    {
                        "code": "HIGH_COMPLAINT_RATE",
                        "message": "卡扣问题严重度达5分",
                        "related_item_id": item1_id,
                        "evidence_refs": ["rev_01"],
                    }
                ],
                model_metadata={"capture": {}},
            )
        )
        session.add(
            Evidence(
                evidence_id="evi_01",
                tenant_id="dev-tenant",
                task_id=task_id,
                item_id=item1_id,
                source_type="REVIEW_TEXT",
                source_ref="rev_01",
                excerpt="The plastic buckle snapped immediately on day two.",
                source_url="https://amazon.com/dp/B0FFWCNZGF",
                source_metadata={"rating": 1},
                provenance={"source": "amazon_reviews"},
            )
        )
        session.commit()
    return task_id, asin1, asin2


def test_export_task_charter_not_found(sqlite_session_factory):
    with sqlite_session_factory() as session:
        with pytest.raises(ApiError) as exc_info:
            export_task_charter_zip(
                session, tenant_id="dev-tenant", task_id="non_existent"
            )
        assert exc_info.value.status_code == 404
        assert exc_info.value.code == "TASK_NOT_FOUND"


def test_export_task_charter_running_raises_409(sqlite_session_factory):
    with sqlite_session_factory() as session:
        session.add(
            Task(
                task_id="tsk_running",
                tenant_id="dev-tenant",
                platform="amazon",
                marketplace="US",
                window_preset="6m",
                status=TaskStatus.RUNNING.value,
            )
        )
        session.commit()

        with pytest.raises(ApiError) as exc_info:
            export_task_charter_zip(
                session, tenant_id="dev-tenant", task_id="tsk_running"
            )
        assert exc_info.value.status_code == 409
        assert exc_info.value.code == "TASK_NOT_READY"


def test_export_task_charter_zip_contents(sqlite_session_factory):
    task_id, asin1, asin2 = _seed_completed_task(sqlite_session_factory)

    with sqlite_session_factory() as session:
        zip_bytes, filename = export_task_charter_zip(
            session,
            tenant_id="dev-tenant",
            task_id=task_id,
        )

    assert filename == f"工程任务书_{task_id}.zip"
    assert len(zip_bytes) > 0

    # Verify ZIP structure
    with zipfile.ZipFile(io.BytesIO(zip_bytes), "r") as zf:
        file_list = zf.namelist()
        assert "工程任务书_建议.docx" in file_list
        assert f"{asin1}.xlsx" in file_list
        assert f"{asin2}.xlsx" in file_list

        # Verify DOCX content
        docx_data = zf.read("工程任务书_建议.docx")
        doc = Document(io.BytesIO(docx_data))
        text_content = "\n".join(p.text for p in doc.paragraphs)
        assert "工程改款任务书" in text_content
        assert "加厚卡扣模具壁厚" in text_content
        assert "加入高密度EPE防震珍珠棉" in text_content
        assert asin1 in text_content
        assert asin2 in text_content

        # Verify XLSX content for item 1
        xlsx1_data = zf.read(f"{asin1}.xlsx")
        wb1 = openpyxl.load_workbook(io.BytesIO(xlsx1_data))
        assert "产品概览" in wb1.sheetnames
        assert "改款建议" in wb1.sheetnames
        assert "用户痛点" in wb1.sheetnames
        assert "原始证据链" in wb1.sheetnames

        ws_prop = wb1["改款建议"]
        prop_rows = list(ws_prop.iter_rows(values_only=True))
        assert len(prop_rows) == 3  # header + 2 proposals
        assert prop_rows[1][2] == "加厚卡扣模具壁厚"
        assert prop_rows[2][2] == "加入高密度EPE防震珍珠棉"

        ws_pain = wb1["用户痛点"]
        pain_rows = list(ws_pain.iter_rows(values_only=True))
        assert len(pain_rows) == 2  # header + 1 pain point
        assert pain_rows[1][1] == "卡扣松动易碎"

        ws_evi = wb1["原始证据链"]
        evi_rows = list(ws_evi.iter_rows(values_only=True))
        assert len(evi_rows) == 2  # header + 1 evidence
        assert "The plastic buckle snapped" in evi_rows[1][3]


@pytest.mark.asyncio
async def test_export_endpoint_returns_zip(sqlite_session_factory):
    import httpx
    from httpx import ASGITransport

    task_id, _, _ = _seed_completed_task(sqlite_session_factory)
    app = create_app()

    def override_get_session():
        with sqlite_session_factory() as session:
            yield session

    app.dependency_overrides[get_session] = override_get_session
    transport = ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.get(f"/api/v1/tasks/{task_id}/export")
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/zip"
        assert "attachment" in resp.headers["content-disposition"]
        assert len(resp.content) > 0
        with zipfile.ZipFile(io.BytesIO(resp.content), "r") as zf:
            assert "工程任务书_建议.docx" in zf.namelist()
